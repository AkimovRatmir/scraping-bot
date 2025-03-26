import os
import asyncio
from playwright.async_api import async_playwright
from openpyxl import load_workbook

# Логин и пароль для входа на сайт Kaspi
login_email = os.environ["KASPI_LOGIN"]
password = os.environ["KASPI_PASSWORD"]

GREEN_CHECK = "\033[92m✔\033[0m"  # Зеленый символ для уведомления
RED_CROSS = "\033[91m✖\033[0m"    # Красный символ для уведомления

excel_path = r"/home/ratmir/kaspiBotBaza/baza.xlsx"

# Функция для записи данных в Excel-файл (обновляет данные, если товар уже существует)
def write_and_save_to_excel(workbook, worksheet, data):
    product_name = data[0]  # Название товара
    found = False
    
    for row in range(2, worksheet.max_row + 1):  # Начинаем со второй строки (пропуская заголовок)
        existing_product = worksheet[f"A{row}"].value
        if existing_product == product_name:
            worksheet[f"B{row}"] = data[1]  # URL
            worksheet[f"C{row}"] = data[2]  # Текущая цена
            worksheet[f"D{row}"] = data[3]  # Цена на странице товара
            found = True
            print(f"Товар '{product_name}' найден, данные обновлены.")
            break
    
    if not found:
        next_row = worksheet.max_row + 1
        worksheet[f"A{next_row}"] = data[0]  # Название товара
        worksheet[f"B{next_row}"] = data[1]  # URL
        worksheet[f"C{next_row}"] = data[2]  # Текущая цена
        worksheet[f"D{next_row}"] = data[3]  # Цена на странице товара
        print(f"Товар '{product_name}' не найден, добавляем новый.")
    
    workbook.save(excel_path)  # Сохранение данных после записи


async def login(page):
    await page.goto('https://kaspi.kz/mc/#/login')
    await page.wait_for_load_state('networkidle')
    await asyncio.sleep(1)

    await page.click('input#user_email_field')
    await asyncio.sleep(1)
    await page.wait_for_selector('input#user_email_field', timeout=10000)
    await page.fill('input#user_email_field', login_email)
    await asyncio.sleep(1)

    await page.click('button.button.is-primary')
    await asyncio.sleep(1)
    
    await page.wait_for_selector('input#password_field', timeout=10000)
    await page.fill('input#password_field', password)
    await asyncio.sleep(1)
    
    await page.click('button.button.is-primary')
    await page.wait_for_load_state('networkidle')
    await asyncio.sleep(2)

    # Переход на страницу активных товаров после логина
    await page.goto('https://kaspi.kz/mc/#/products/active/1')
    await page.wait_for_load_state('networkidle')
    await asyncio.sleep(2)

async def get_product_url(page, product_name):
    try:
        print(f"Ищем продукт с именем: {product_name}")
        await page.click(f"//a[contains(text(), '{product_name}')]")
        await page.wait_for_load_state('networkidle')
        await asyncio.sleep(2)

        locator = page.locator(
            '#app > section > div.main-content > div > div.master-product__content > div.master-product-info > div > div:nth-child(2) > div:nth-child(2) > a'
        )

        if await locator.count() > 0:
            product_url = await locator.get_attribute('href')
            print(f"Найден URL для товара '{product_name}': {product_url}")
            return product_url
        else:
            print(f"Не удалось найти URL для товара: {product_name}")
            return None

    except Exception as e:
        print(f"Ошибка при получении URL для товара '{product_name}': {e}")
        return None

async def analyze_and_update_price(page, product_name, product_url, current_price):
    try:
        full_url = f"https://kaspi.kz{product_url}"
        await page.goto(full_url)
        await page.wait_for_load_state('networkidle')
        await asyncio.sleep(2)

        # Получаем текущую цену товара на странице
        price_selector = '#ItemView > div.item > div > div.item__inner-right > div > div:nth-child(1) > div.item__price > div.item__price-left-side > div.item__price-once'
        competitor_price_element = await page.query_selector(price_selector)
        competitor_price_text = await competitor_price_element.inner_text()
        competitor_price = float(competitor_price_text.replace('₸', '').replace(' ', '').replace(' ', '').strip())

        # Получаем цены всех продавцов на странице товара
        sellers_prices_elements = await page.query_selector_all('div.sellers-table__price-cell-text:not(._installments-price)')
        sellers_prices = [
            float((await element.inner_text()).replace('₸', '').replace(' ', '').replace(' ', '').strip())
            for element in sellers_prices_elements
        ]

        # Минимальная цена среди других продавцов, исключая текущую цену
        other_sellers_prices = [price for price in sellers_prices if price != current_price]
        min_seller_price = min(other_sellers_prices, default=None)

        print(f"Текущая цена: {current_price} ₸, Цена на странице товара: {competitor_price} ₸, Минимальная цена среди других продавцов: {min_seller_price} ₸")

        # Если минимальная цена конкурентов ниже текущей цены, устанавливаем новую цену
        if min_seller_price is not None and min_seller_price < current_price:
            new_price = min_seller_price - 10
            print(f"{GREEN_CHECK} Новая цена для товара: {new_price} ₸")

            # Переход к странице с оффером товара на сайте
            await page.goto('https://kaspi.kz/mc/#/products/active/1')
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)

            # Переход на оффер товара, используя имя товара
            await page.click(f"//a[contains(text(), '{product_name}')]")
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)

            # Нажимаем на кнопку "Изменить цену"
            await page.click("a:has-text('Изменить цену')")
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)

            # Ввод новой цены
            input_selector = '#app > section > div.main-content > div > div.block > div.b-table > div > table > thead > tr.is-subheading > th:nth-child(2) > div > span > div:nth-child(2) > input'

            # Ожидание и взаимодействие с элементом
            await page.wait_for_selector(input_selector, timeout=10000)  # Ожидание появления элемента
            await page.fill(input_selector, 'Ваш текст или значение')  # Заполнение поля ввода

            # Очищаем поле и вводим новую цену
            await page.click(input_selector)
            await page.keyboard.press("Control+A")
            await page.keyboard.press("Backspace")
            await asyncio.sleep(1)
            await page.fill(input_selector, str(int(new_price)))

            # Нажимаем кнопку "Сохранить"
            save_button_selector = '#app > section > div.main-content > div > div.block > div.block > button.button.button.is-primary.mr-1 > span'
            await page.wait_for_selector(save_button_selector, timeout=10000)
            await page.click(save_button_selector)
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(1)

            print(f'Цена успешно обновлена до {new_price} ₸')
        
        else:
            print(f"{GREEN_CHECK} Изменение цены не требуется: текущая цена уже минимальная или равна конкурентной.")
        
        return competitor_price, min_seller_price

    except Exception as e:
        print(f"Ошибка при анализе и изменении цены для товара '{product_name}': {e}")
        return None, None

async def get_all_product_info(page, workbook, worksheet):
    current_page = 1
    while True:
        await page.goto(f'https://kaspi.kz/mc/#/products/active/{current_page}')
        await page.wait_for_load_state('networkidle')
        await asyncio.sleep(2)

        current_page_products = await page.evaluate('''() => {
            return Array.from(document.querySelectorAll('tr')).map(row => {
                const productNameElement = row.querySelector('td[data-label="Товар"] p.is-5 a');
                const productName = productNameElement ? productNameElement.innerText.trim() : null;
                const productPriceElement = row.querySelector('td[data-label="Цена, тенге"] p.subtitle.is-5');
                const productPrice = productPriceElement ? parseFloat(productPriceElement.innerText.replace('₸', '').replace(' ', '').replace(' ', '').trim()) : null;
                if (productName && productPrice !== null) {
                    return { productName, productPrice };
                } else {
                    return null;
                }
            }).filter(product => product !== null);
        }''')

        if not current_page_products:
            print("На текущей странице нет товаров.")
            break
        
        for idx, product in enumerate(current_page_products, start=2):
            product_url = await get_product_url(page, product['productName'])
            if product_url:
                competitor_price, min_seller_price = await analyze_and_update_price(page, product['productName'], product_url, product['productPrice'])
                if competitor_price is not None:
                    data = [product['productName'], f"https://kaspi.kz{product_url}", product['productPrice'], competitor_price]
                    print(f"Записываем данные в Excel: {data}")
                    write_and_save_to_excel(workbook, worksheet, data)
            else:
                print(f"Не удалось получить URL для товара: {product['productName']}")

            await page.goto(f'https://kaspi.kz/mc/#/products/active/{current_page}')
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)

        print(f"Текущая страница: обработано {len(current_page_products)} товаров.")

        next_button = await page.query_selector('#offers-table > div.level-right > div > nav > a.pagination-link.pagination-next')
        if next_button:
            current_page += 1
            await next_button.click()
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)
        else:
            print("Все страницы обработаны.")
            break

async def process_products(playwright):
    print("Логинимся на сайт Kaspi.kz...")
    browser = await playwright.chromium.launch(headless=False)
    page = await browser.new_page()

    await login(page)

    print("Переход на страницу активных товаров...")
    
    print(f"Открываем файл Excel: {excel_path}")
    workbook = load_workbook(excel_path)
    worksheet = workbook.active

    await get_all_product_info(page, workbook, worksheet)

    await browser.close()

async def main():
    async with async_playwright() as playwright:
        await process_products(playwright)

# Запуск основной функции
asyncio.run(main())