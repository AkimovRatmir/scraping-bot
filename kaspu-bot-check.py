import os
import asyncio
from playwright.async_api import async_playwright
from openpyxl import load_workbook

# Логин и пароль для входа на сайт Kaspi
login_email = os.environ["KASPI_LOGIN"]
password = os.environ["KASPI_PASSWORD"]

GREEN_CHECK = "\033[92m✔\033[0m"  # Зеленый символ
RED_CROSS = "\033[91m✖\033[0m"    # Красный символ

excel_path = r"/home/ratmir/kaspiBotBaza/baza.xlsx"

# Функция для записи данных в Excel-файл
def write_and_save_to_excel(workbook, worksheet, data):
    product_name = data[0]  
    found = False
    
    for row in range(2, worksheet.max_row + 1):  
        existing_product = worksheet[f"A{row}"].value
        if existing_product == product_name:
            worksheet[f"B{row}"] = data[1]  
            worksheet[f"C{row}"] = data[2]  
            worksheet[f"D{row}"] = data[3]  
            worksheet[f"E{row}"] = data[4]  
            worksheet[f"F{row}"] = data[5]  
            found = True
            print(f"✅ Товар '{product_name}' найден, данные обновлены.")
            break
    
    if not found:
        next_row = worksheet.max_row + 1
        worksheet[f"A{next_row}"] = data[0]  
        worksheet[f"B{next_row}"] = data[1]  
        worksheet[f"C{next_row}"] = data[2]  
        worksheet[f"D{next_row}"] = data[3]  
        worksheet[f"E{next_row}"] = data[4]  
        worksheet[f"F{next_row}"] = data[5]  
        worksheet[f"H{next_row}"] = 1  
        worksheet[f"I{next_row}"] = 1  
        print(f"✅ Товар '{product_name}' не найден, добавляем новый.")

    workbook.save(excel_path)  

# Функция логина
async def login(page):
    await page.goto('https://kaspi.kz/mc/#/login')
    await page.wait_for_load_state('networkidle')
    await asyncio.sleep(1)

    await page.fill('input#user_email_field', login_email)
    await asyncio.sleep(1)

    await page.click('button.button.is-primary')
    await page.wait_for_selector('input#password_field', timeout=10000)
    await page.fill('input#password_field', password)
    await asyncio.sleep(1)
    
    await page.click('button.button.is-primary')
    await page.wait_for_load_state('networkidle')
    await asyncio.sleep(2)

    await page.goto('https://kaspi.kz/mc/#/products/active/1')
    await page.wait_for_load_state('networkidle')
    await asyncio.sleep(2)

# Функция для получения URL, SKU и изображения товара
async def get_product_url_and_sku(page, product_name):
    try:
        print(f"🔎 Ищем продукт: {product_name}")

        # Поиск товара по имени
        product_locator = page.locator(f"text={product_name}")
        if await product_locator.count() == 0:
            print(f"{RED_CROSS} Товар '{product_name}' не найден на странице.")
            return None, None, None

        await product_locator.first.click()
        await page.wait_for_load_state('networkidle')
        await asyncio.sleep(2)

        # Получение URL товара
        product_url_locator = page.locator('a.master-product-info__link')
        product_url = await product_url_locator.get_attribute('href') if await product_url_locator.count() > 0 else None

        # Получение SKU товара
        sku_locator = page.locator('div.master-product-info__sku')
        product_sku = await sku_locator.inner_text() if await sku_locator.count() > 0 else None

        # Получение изображения товара
        image_locator = page.locator('div.thumbnail-container img')
        product_image_url = await image_locator.get_attribute('src') if await image_locator.count() > 0 else None

        if product_url:
            print(f"{GREEN_CHECK} Найден URL: {product_url}")
        else:
            print(f"{RED_CROSS} URL не найден для {product_name}")

        return product_url, product_sku, product_image_url

    except Exception as e:
        print(f"{RED_CROSS} Ошибка при получении данных для '{product_name}': {e}")
        return None, None, None

# Функция для получения информации о товарах
async def get_all_product_info(page, workbook, worksheet):
    current_page = 1
    while True:
        await page.goto(f'https://kaspi.kz/mc/#/products/active/{current_page}')
        await page.wait_for_load_state('networkidle')
        await asyncio.sleep(2)

        # Получаем список всех товаров на странице
        current_page_products = await page.evaluate('''() => {
            return Array.from(document.querySelectorAll('tr')).map(row => {
                const productNameElement = row.querySelector('td[data-label="Товар"] p.is-5 a');
                const productName = productNameElement ? productNameElement.innerText.trim() : null;
                const productPriceElement = row.querySelector('td[data-label="Цена, тенге"] p.subtitle.is-5');
                const productPrice = productPriceElement ? parseFloat(productPriceElement.innerText.replace('₸', '').replace(' ', '').replace(' ', '').trim()) : null;
                if (productName && productPrice !== null) {
                    return { productName, productPrice };
                } else {
                    return null;
                }
            }).filter(product => product !== null);
        }''')

        if not current_page_products:
            print("❌ На текущей странице нет товаров.")
            break
        
        for idx, product in enumerate(current_page_products, start=2):
            try:
                product_url, product_sku, product_image_url = await get_product_url_and_sku(page, product['productName'])
                if not product_url:
                    print(f"{RED_CROSS} Пропускаем {product['productName']} из-за отсутствия URL")
                    continue  

                data = [product['productName'], f"https://kaspi.kz{product_url}", product['productPrice'], None, product_sku, product_image_url]
                print(f"✅ Записываем данные в Excel: {data}")
                write_and_save_to_excel(workbook, worksheet, data)

            except Exception as e:
                print(f"❌ Ошибка при обработке товара {product['productName']}: {e}")
                continue

            await page.goto(f'https://kaspi.kz/mc/#/products/active/{current_page}')
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)

        next_button = await page.query_selector('a.pagination-next')
        if next_button:
            current_page += 1
            await next_button.click()
            await page.wait_for_load_state('networkidle')
            await asyncio.sleep(2)
        else:
            print("✅ Все страницы обработаны.")
            break

async def process_products(playwright):
    print("🔐 Логинимся на сайт Kaspi.kz...")
    browser = await playwright.chromium.launch(headless=False)
    page = await browser.new_page()

    await login(page)

    print("📂 Открываем Excel: ", excel_path)
    workbook = load_workbook(excel_path)
    worksheet = workbook.active

    await get_all_product_info(page, workbook, worksheet)

    await browser.close()

async def main():
    async with async_playwright() as playwright:
        await process_products(playwright)

asyncio.run(main())
